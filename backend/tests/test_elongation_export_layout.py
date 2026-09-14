"""Sanitized structural regressions: two Excel rows per physical measurement."""

from copy import deepcopy
from decimal import Decimal
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from app.services.elongations.template import analyse_template, build_export_xlsx


def field_template() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Operativa"
    sheet["A1"] = "OBRA: EJEMPLO"
    for letter, width in {
        "A": 11,
        "B": 15,
        "C": 14,
        "D": 15,
        "H": 15,
        "I": 10,
        "J": 15,
        "K": 11,
    }.items():
        sheet.column_dimensions[letter].width = width
    for letter in "EFG":
        sheet.column_dimensions[letter].hidden = True
    side = Side(style="thin", color="000000")
    for start, label, title in ((3, "T321", "BANDAS"), (11, "T300", "DISTRIBUIDOS")):
        sheet.merge_cells(start_row=start, end_row=start + 1, start_column=1, end_column=11)
        sheet.cell(start, 1, title)
        header, first = start + 2, start + 4
        for col, text in {
            1: "Item",
            2: "Label",
            3: "Longitud (m)",
            4: "Cantidad de tendones",
            8: "Elongación (cm)",
        }.items():
            sheet.cell(header, col, text)
            if col < 8:
                sheet.merge_cells(
                    start_row=header, end_row=header + 1, start_column=col, end_column=col
                )
        sheet.merge_cells(start_row=header, end_row=header, start_column=8, end_column=11)
        for col, text in {8: "Calculada", 9: "Max.", 10: "Elong. Medida", 11: "Min."}.items():
            sheet.cell(header + 1, col, text)
        for row in range(first, first + 4):
            sheet.row_dimensions[row].height = 11.4 if row == first else 12
            for col in (1, 2, 3, 4, 8, 9, 10, 11):
                cell = sheet.cell(row, col)
                cell.font = Font(name="Arial", size=12)
                cell.border = Border(left=side, right=side, top=side, bottom=side)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.number_format = "0.000" if col in (3, 8, 9, 10, 11) else "General"
        for row in (first, first + 2):
            for col, value in {
                1: 1,
                8: 7,
                9: f"=H{row}+(H{row}*0.07)",
                10: 99,
                11: f"=H{row}-(H{row}*0.07)",
            }.items():
                sheet.cell(row, col, value)
                sheet.merge_cells(start_row=row, end_row=row + 1, start_column=col, end_column=col)
        for col, value in {2: label, 3: 10, 4: 2}.items():
            sheet.cell(first, col, value)
            sheet.merge_cells(start_row=first, end_row=first + 3, start_column=col, end_column=col)
    sheet.conditional_formatting.add(
        "J7:J18", FormulaRule(formula=["J7>I999"], fill=PatternFill("solid", fgColor="FF0000"))
    )
    workbook.create_sheet("Notas")["A1"] = "Conservar esta hoja"
    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


@pytest.mark.parametrize("final", [False, True])
def test_export_orders_all_labels_and_preserves_complete_physical_row_height(final: bool) -> None:
    template = field_template()
    mapping = analyse_template(template)
    # No label matches the example. Classes, arrival order and stale numeric
    # metadata must not partition/reorder the report or change approved data.
    groups = []
    for label, count, classification in (
        (203, 1, "band"),
        (210, 2, "band"),
        (200, 2, "distributed"),
        (229, 3, "band"),
        (202, 1, "distributed"),
        (10, 1, "band"),
        (2, 1, "distributed"),
    ):
        groups.append(
            {
                "label": f"T{label}",
                "label_number": 999 - label,
                "classification": classification,
                "length_m": Decimal("7.000") + Decimal(label) / 1000,
                "strand_count": count,
                "calculated_elongation": Decimal("5.000") + Decimal(label) / 1000,
                "measurements": [
                    {
                        "ordinal": ordinal,
                        "measured_elongation": Decimal("5.000") + Decimal(label + ordinal) / 1000,
                        "review_status": "approved",
                        "tolerance_status": "within",
                    }
                    for ordinal in range(1, count + 1)
                ],
            }
        )
    before = deepcopy(groups)
    result = build_export_xlsx(
        template,
        mapping,
        groups,
        final=final,
        history={"kind": "final" if final else "theoretical"},
    )
    workbook = load_workbook(BytesIO(result))
    sheet = workbook["Operativa"]
    label_rows = [
        (row, sheet.cell(row, 2).value)
        for row in range(7, sheet.max_row + 1)
        if sheet.cell(row, 2).value is not None
    ]
    assert [label for _, label in label_rows] == [
        "T2",
        "T10",
        "T200",
        "T202",
        "T203",
        "T210",
        "T229",
    ]
    assert groups == before
    assert (
        sum(isinstance(sheet.cell(row, 1).value, int) for row in range(7, sheet.max_row + 1)) == 11
    )
    assert [sheet.cell(row, 1).value for row in range(7, 18)] == list(range(1, 12))
    for row, label in label_rows:
        group = next(group for group in groups if group["label"] == label)
        count = group["strand_count"]
        assert sheet.cell(row, 12).value == (
            "Banda" if group["classification"] == "band" else "Distribuido"
        )
        assert Decimal(str(sheet.cell(row, 3).value)) == group["length_m"]
        assert sheet.cell(row, 4).value == count
        if count > 1:
            for col in ("B", "C", "D", "L"):
                assert f"{col}{row}:{col}{row + count - 1}" in sheet.merged_cells
                assert sheet[f"{col}{row + count - 1}"].border.bottom.style == "thin"
        for ordinal in range(1, count + 1):
            physical_row = row + ordinal - 1
            assert Decimal(str(sheet.cell(physical_row, 8).value)) == group["calculated_elongation"]
            measurement = sheet.cell(physical_row, 10).value
            if final:
                assert (
                    Decimal(str(measurement))
                    == group["measurements"][ordinal - 1]["measured_elongation"]
                )
            else:
                assert measurement is None
    for row in range(7, 18):
        assert sheet.row_dimensions[row].height == pytest.approx(23.4)
        assert sheet.cell(row, 9).value == f"=H{row}+(H{row}*0.07)"
        assert sheet.cell(row, 11).value == f"=H{row}-(H{row}*0.07)"
        assert sheet.cell(row, 9).border.bottom.style == "thin"
    assert "A3:L4" in sheet.merged_cells
    assert "H5:K5" in sheet.merged_cells
    assert "L5:L6" in sheet.merged_cells
    assert sheet["A3"].value == "TENDONES POR LABEL"
    assert sheet["L5"].value == "Clase"
    assert sheet.column_dimensions["B"].width == 15
    assert sheet.column_dimensions["E"].hidden
    assert sheet.print_area == "'Operativa'!$A$1:$L$17"
    assert sheet.print_title_rows == "$5:$6"
    assert sheet.page_setup.fitToWidth == 1
    assert sheet.page_setup.fitToHeight == 0
    assert len(sheet.conditional_formatting) == 1
    assert workbook["Notas"]["A1"].value == "Conservar esta hoja"


def test_manual_print_breaks_do_not_cut_a_group_of_physical_tendons() -> None:
    template = field_template()
    mapping = analyse_template(template)
    groups = [
        {
            "label": f"T{100 + index}",
            "classification": "band" if index % 2 else "distributed",
            "strand_count": 8,
            "length_m": Decimal("10"),
            "calculated_elongation": Decimal("7"),
            "measurements": [],
        }
        for index in range(7)
    ]
    result = build_export_xlsx(template, mapping, groups, final=False, history={})
    worksheet = load_workbook(BytesIO(result))["Operativa"]
    starts = {row for row in range(7, worksheet.max_row + 1) if worksheet.cell(row, 2).value}
    assert len(worksheet.row_breaks.brk) >= 2
    for page_break in worksheet.row_breaks.brk:
        assert page_break.id + 1 in starts
        assert page_break.man
